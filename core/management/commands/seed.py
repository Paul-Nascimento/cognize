"""
Popula o banco com dados iniciais:
- 4 usuários (um por cargo), senha padrão 'cognize123'
- Empresas da planilha de cadastro (apenas ID + razão social; consulta fica pendente)
- As tarefas do 'Relatório'

Uso: python manage.py seed
"""
import re
import unicodedata
from pathlib import Path

from django.core.management.base import BaseCommand
from django.db import transaction

from accounts.models import Cargo, User
from empresas.models import Empresa, RegimeTributario, StatusConsulta
from tarefas.models import Departamento, Orgao, Tarefa


def _norm(t):
    if t is None:
        return ""
    t = str(t).strip().lower()
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")


def _match(valor, choices, default):
    n = _norm(valor)
    for value, label in choices:
        if _norm(value) == n or _norm(label) == n:
            return value
    return default


class Command(BaseCommand):
    help = "Popula o banco com usuários, empresas e tarefas de exemplo."

    def add_arguments(self, parser):
        parser.add_argument(
            "--empresas", type=str, default="",
            help="Caminho da planilha de empresas (opcional).",
        )
        parser.add_argument(
            "--tarefas", type=str, default="",
            help="Caminho da planilha de tarefas/Relatório (opcional).",
        )

    @transaction.atomic
    def handle(self, *args, **opts):
        self._criar_usuarios()
        if opts["empresas"] and Path(opts["empresas"]).exists():
            self._importar_empresas(opts["empresas"])
        if opts["tarefas"] and Path(opts["tarefas"]).exists():
            self._importar_tarefas(opts["tarefas"])
        self.stdout.write(self.style.SUCCESS("Seed concluído."))

    def _criar_usuarios(self):
        base = [
            ("admin", "Ana", "Administradora", Cargo.ADMINISTRADOR, True),
            ("gestor", "Gustavo", "Gestor", Cargo.GESTOR, False),
            ("colab", "Carla", "Colaboradora", Cargo.COLABORADOR, False),
            ("visual", "Victor", "Visualizador", Cargo.VISUALIZADOR, False),
        ]
        for username, first, last, cargo, is_super in base:
            user, criado = User.objects.get_or_create(
                username=username,
                defaults={
                    "first_name": first, "last_name": last, "cargo": cargo,
                    "is_staff": is_super, "is_superuser": is_super,
                    "email": f"{username}@cognize.local",
                },
            )
            if criado:
                user.set_password("cognize123")
                user.save()
                self.stdout.write(f"  usuário '{username}' criado ({cargo}).")

    def _importar_empresas(self, caminho):
        from openpyxl import load_workbook

        wb = load_workbook(caminho, data_only=True, read_only=True)
        ws = wb["Sheet"] if "Sheet" in wb.sheetnames else wb.active
        n = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] in (None, ""):
                continue
            id_sistema = str(row[0]).strip()
            if re.fullmatch(r"\d+\.0", id_sistema):
                id_sistema = str(int(float(id_sistema)))
            razao = str(row[1] or "").strip()
            # CNPJ fictício determinístico (planilha não traz CNPJ real)
            base = re.sub(r"\D", "", id_sistema).zfill(8)[:8]
            cnpj = f"{base[:2]}{base[2:5]}{base[5:8]}000100"[:14].ljust(14, "0")
            Empresa.objects.get_or_create(
                id_sistema=id_sistema,
                defaults={
                    "cnpj": cnpj,
                    "razao_social": razao,
                    "regime_tributario": RegimeTributario.SIMPLES,
                    "status_consulta": StatusConsulta.PENDENTE,
                    "mensagem_consulta": "Cadastrada via seed — consulta pendente.",
                },
            )
            n += 1
        self.stdout.write(f"  {n} empresas processadas.")

    def _importar_tarefas(self, caminho):
        from openpyxl import load_workbook

        wb = load_workbook(caminho, data_only=True, read_only=True)
        ws = wb.active
        linhas = list(ws.iter_rows(values_only=True))
        headers = [_norm(h) for h in linhas[0]]

        def col(nome):
            for i, h in enumerate(headers):
                if nome in h:
                    return i
            return None

        c_dep, c_nome = col("departamento"), col("nome")
        c_freq, c_acao = col("frequenc"), col("acao")
        c_meta, c_dia = col("meta"), col("vencimento")
        c_fato = col("fato gerador")
        c_org, c_status = col("orgao"), None
        for i, h in enumerate(headers):
            if h == "status":
                c_status = i

        n = 0
        for row in linhas[1:]:
            if c_nome is None or c_nome >= len(row):
                continue
            nome = str(row[c_nome] or "").strip()
            nome = nome.replace("-->", "").replace("*", "").strip(" |")
            if not nome:
                continue
            dia = None
            if c_dia is not None and row[c_dia] not in (None, "", 0):
                try:
                    dia = int(float(row[c_dia])) or None
                except (TypeError, ValueError):
                    dia = None
            Tarefa.objects.get_or_create(
                nome=nome,
                departamento=_match(
                    row[c_dep] if c_dep is not None else "", Departamento.choices, Departamento.OUTRO
                ),
                defaults={
                    "frequencia": str(row[c_freq] or "Mensal").strip() if c_freq is not None else "Mensal",
                    "acao_qtd_dias": str(row[c_acao] or "").strip() if c_acao is not None else "",
                    "meta_qtd_dias": str(row[c_meta] or "").strip() if c_meta is not None else "",
                    "dia_vencimento": dia,
                    "fato_gerador": str(row[c_fato] or "").strip() if c_fato is not None else "",
                    "orgao": _match(row[c_org] if c_org is not None else "", Orgao.choices, Orgao.FEDERAL),
                    "ativo": True,
                },
            )
            n += 1
        self.stdout.write(f"  {n} tarefas processadas.")
