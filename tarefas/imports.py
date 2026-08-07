"""Importação de tarefas e de vínculos empresa-tarefa via planilha."""
import io
import unicodedata

from openpyxl import Workbook, load_workbook

from empresas.models import Empresa
from .models import Departamento, Orgao, Tarefa, VinculoTarefa


def _norm(texto) -> str:
    if texto is None:
        return ""
    txt = str(texto).strip().lower()
    return "".join(
        c for c in unicodedata.normalize("NFD", txt) if unicodedata.category(c) != "Mn"
    )


def _match_choice(valor, choices, default):
    n = _norm(valor)
    for value, label in choices:
        if _norm(value) == n or _norm(label) == n:
            return value
    return default


def _mapear(headers):
    idx = {}
    for i, h in enumerate(headers):
        n = _norm(h)
        if n.startswith("departamento"):
            idx["departamento"] = i
        elif n == "nome" or n.startswith("nome"):
            idx.setdefault("nome", i)
        elif n.startswith("frequenc"):
            idx["frequencia"] = i
        elif "acao" in n:
            idx["acao"] = i
        elif "meta" in n:
            idx["meta"] = i
        elif "vencimento" in n:
            idx["dia"] = i
        elif "fato gerador" in n or "competencia" in n:
            idx["fato"] = i
        elif n.startswith("orgao"):
            idx["orgao"] = i
        elif n.startswith("re-envio") or n.startswith("reenvio"):
            idx["reenvio"] = i
        elif n == "status":
            idx["status"] = i
    return idx


def importar_tarefas(arquivo):
    wb = load_workbook(arquivo, data_only=True, read_only=True)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return {"criadas": 0, "atualizadas": 0, "erros": ["Planilha vazia."]}

    idx = _mapear(linhas[0])
    if "nome" not in idx:
        return {"criadas": 0, "atualizadas": 0, "erros": ["Coluna 'Nome' não encontrada."]}

    criadas = atualizadas = 0
    erros = []
    for n, linha in enumerate(linhas[1:], start=2):
        def val(chave):
            i = idx.get(chave)
            return linha[i] if i is not None and i < len(linha) else None

        nome = str(val("nome") or "").strip()
        # limpa marcadores como " --> " e " * " presentes no relatório original
        nome = nome.replace("-->", "").replace(" \\*", "").replace("*", "").strip(" |")
        if not nome:
            continue

        dia_raw = val("dia")
        try:
            dia = int(float(dia_raw)) if dia_raw not in (None, "", 0, "0") else None
            if dia == 0:
                dia = None
        except (TypeError, ValueError):
            dia = None

        defaults = {
            "departamento": _match_choice(
                val("departamento"), Departamento.choices, Departamento.OUTRO
            ),
            "frequencia": str(val("frequencia") or "Mensal").strip() or "Mensal",
            "acao_qtd_dias": str(val("acao") or "").strip(),
            "meta_qtd_dias": str(val("meta") or "").strip(),
            "dia_vencimento": dia,
            "fato_gerador": str(val("fato") or "").strip(),
            "orgao": _match_choice(val("orgao"), Orgao.choices, Orgao.FEDERAL),
            "reenvio": str(val("reenvio") or "").strip(),
            "ativo": _norm(val("status")) in ("", "ativo", "ativa", "sim", "true"),
        }
        try:
            obj, criado = Tarefa.objects.update_or_create(
                nome=nome, departamento=defaults["departamento"], defaults=defaults
            )
            criadas += int(criado)
            atualizadas += int(not criado)
        except Exception as exc:  # noqa: BLE001
            erros.append(f"Linha {n} ('{nome}'): {exc}")

    return {"criadas": criadas, "atualizadas": atualizadas, "erros": erros}


def importar_vinculos(arquivo):
    """Planilha com colunas: ID Empresa | Tarefa (nome)."""
    wb = load_workbook(arquivo, data_only=True, read_only=True)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return {"criados": 0, "erros": ["Planilha vazia."]}

    headers = [_norm(h) for h in linhas[0]]
    idx_emp = idx_tar = None
    for i, h in enumerate(headers):
        if "empresa" in h or ("id" in h and "sistema" in h):
            idx_emp = i
        elif "tarefa" in h:
            idx_tar = i
    if idx_emp is None or idx_tar is None:
        return {"criados": 0, "erros": ["Preciso das colunas 'ID Empresa' e 'Tarefa'."]}

    criados = 0
    erros = []
    for n, linha in enumerate(linhas[1:], start=2):
        id_emp = linha[idx_emp] if idx_emp < len(linha) else None
        nome_tar = linha[idx_tar] if idx_tar < len(linha) else None
        if not id_emp and not nome_tar:
            continue
        try:
            id_emp = str(int(float(id_emp))) if str(id_emp).replace(".", "").isdigit() else str(id_emp).strip()
        except (TypeError, ValueError):
            id_emp = str(id_emp).strip()
        nome_tar = str(nome_tar or "").strip()

        empresa = Empresa.objects.filter(id_sistema=id_emp).first()
        if not empresa:
            erros.append(f"Linha {n}: empresa '{id_emp}' não encontrada.")
            continue
        tarefa = Tarefa.objects.filter(nome__iexact=nome_tar).first()
        if not tarefa:
            erros.append(f"Linha {n}: tarefa '{nome_tar}' não encontrada.")
            continue
        _, criado = VinculoTarefa.objects.get_or_create(empresa=empresa, tarefa=tarefa)
        criados += int(criado)

    return {"criados": criados, "erros": erros}


def gerar_modelo_tarefas() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tarefas"
    headers = [
        "Departamento", "Nome", "Frequência", "Ação(qtd dias)",
        "Meta(qtd dias)", "Dia Vencimento", "Fato Gerador", "Órgão", "Status",
    ]
    ws.append(headers)
    ws.append([
        "Fiscal", "DAS SIMPLES", "Mensal", "3", "4", "20",
        "Mês anterior", "Federal", "Ativo",
    ])
    for col in ws.columns:
        largura = max(len(str(c.value or "")) for c in col) + 3
        ws.column_dimensions[col[0].column_letter].width = min(largura, 40)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_modelo_vinculos() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vínculos"
    ws.append(["ID Empresa", "Tarefa"])
    ws.append(["300", "DAS SIMPLES"])
    ws.append(["300", "Folha de pagamento"])
    ws.append(["401", "DAS SIMPLES"])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
