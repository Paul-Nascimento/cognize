"""Importação de empresas a partir de planilha .xlsx e geração de modelo."""
import io
import re
import unicodedata

from openpyxl import Workbook, load_workbook

from .models import Empresa, RegimeTributario, somente_digitos

# Mapeia rótulos de regime aceitos na planilha -> valor interno
_REGIMES = {
    "simples": RegimeTributario.SIMPLES,
    "simples nacional": RegimeTributario.SIMPLES,
    "mei": RegimeTributario.MEI,
    "presumido": RegimeTributario.PRESUMIDO,
    "lucro presumido": RegimeTributario.PRESUMIDO,
    "real": RegimeTributario.REAL,
    "lucro real": RegimeTributario.REAL,
    "imune": RegimeTributario.IMUNE,
    "isenta": RegimeTributario.IMUNE,
    "imune / isenta": RegimeTributario.IMUNE,
}


def _norm(texto) -> str:
    if texto is None:
        return ""
    txt = str(texto).strip().lower()
    txt = "".join(
        c for c in unicodedata.normalize("NFD", txt) if unicodedata.category(c) != "Mn"
    )
    return txt


def _mapear_cabecalho(headers):
    """Localiza índices das colunas por nome, tolerante a variações."""
    idx = {}
    for i, h in enumerate(headers):
        n = _norm(h)
        if "id" in n and "sistema" in n:
            idx["id_sistema"] = i
        elif n in ("id", "codigo", "empresa"):
            idx.setdefault("id_sistema", i)
        elif "cnpj" in n:
            idx["cnpj"] = i
        elif "regime" in n:
            idx["regime"] = i
        elif "iss" in n:
            idx["iss"] = i
        elif "icms" in n:
            idx["icms"] = i
        elif "observ" in n:
            idx["obs"] = i
    return idx


def importar_empresas(arquivo, usuario=None):
    """
    Lê a planilha e cria/atualiza empresas.
    Retorna dict com 'criadas', 'atualizadas', 'erros' (lista de strings).
    """
    wb = load_workbook(arquivo, data_only=True, read_only=True)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return {"criadas": 0, "atualizadas": 0, "erros": ["Planilha vazia."]}

    idx = _mapear_cabecalho(linhas[0])
    if "id_sistema" not in idx or "cnpj" not in idx:
        return {
            "criadas": 0, "atualizadas": 0,
            "erros": ["Não encontrei as colunas obrigatórias 'ID Sistema' e 'CNPJ'."],
        }

    criadas = atualizadas = 0
    erros = []
    for n, linha in enumerate(linhas[1:], start=2):
        def val(chave):
            i = idx.get(chave)
            return linha[i] if i is not None and i < len(linha) else None

        id_sistema = val("id_sistema")
        cnpj_raw = val("cnpj")
        if id_sistema in (None, "") and cnpj_raw in (None, ""):
            continue  # linha em branco

        id_sistema = str(id_sistema).strip().rstrip(".0") if id_sistema else ""
        # normaliza casos como "300.0" vindos do Excel
        if re.fullmatch(r"\d+\.0", str(val("id_sistema") or "")):
            id_sistema = str(int(float(val("id_sistema"))))

        cnpj = somente_digitos(str(cnpj_raw or ""))
        if len(cnpj) != 14:
            erros.append(f"Linha {n}: CNPJ inválido ('{cnpj_raw}').")
            continue

        regime = _REGIMES.get(_norm(val("regime")), RegimeTributario.SIMPLES)
        defaults = {
            "cnpj": cnpj,
            "regime_tributario": regime,
            "inscricao_iss": str(val("iss") or "").strip(),
            "inscricao_icms": str(val("icms") or "").strip(),
            "observacoes": str(val("obs") or "").strip(),
        }
        try:
            obj, criado = Empresa.objects.update_or_create(
                id_sistema=id_sistema, defaults=defaults
            )
            if criado:
                if usuario:
                    obj.criado_por = usuario
                    obj.save(update_fields=["criado_por"])
                criadas += 1
            else:
                atualizadas += 1
        except Exception as exc:  # noqa: BLE001
            erros.append(f"Linha {n} (ID {id_sistema}): {exc}")

    return {"criadas": criadas, "atualizadas": atualizadas, "erros": erros}


def gerar_modelo_empresas() -> bytes:
    """Gera um .xlsx modelo para preenchimento."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    headers = [
        "ID Sistema", "CNPJ", "Regime Tributário",
        "Inscrição ISS", "Inscrição ICMS", "Observações",
    ]
    ws.append(headers)
    ws.append([
        "300", "00.000.000/0001-00", "Simples Nacional",
        "123456", "", "Empresa exemplo — apague esta linha",
    ])
    ws.append([
        "401", "11.111.111/0001-11", "Lucro Presumido", "", "987654", "",
    ])
    for col in ws.columns:
        largura = max(len(str(c.value or "")) for c in col) + 4
        ws.column_dimensions[col[0].column_letter].width = min(largura, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
